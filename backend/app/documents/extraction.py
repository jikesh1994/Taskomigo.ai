"""Identify uploaded documents by their bytes and extract plain text from them.

The declared filename and content type are never trusted: a file is a PDF only if it
starts with the PDF signature, and a DOCX only if it's a ZIP containing a Word body.
"""

from __future__ import annotations

import io
import re
import unicodedata
import zipfile
from dataclasses import dataclass
from enum import StrEnum

MAX_TEXT_CHARS = 100_000
# DOCX is a ZIP; refuse archives that would inflate beyond this (zip-bomb guard).
MAX_DOCX_UNCOMPRESSED = 50 * 1024 * 1024
MAX_DOCX_ENTRIES = 2_000


class DocumentType(StrEnum):
    PDF = "pdf"
    DOCX = "docx"


CONTENT_TYPES = {
    DocumentType.PDF: "application/pdf",
    DocumentType.DOCX: "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
}


class DocumentError(Exception):
    """The document is unsupported, damaged or unreadable. `message` is user-safe."""

    def __init__(self, message: str, *, code: str) -> None:
        super().__init__(message)
        self.message = message
        self.code = code


@dataclass(frozen=True, slots=True)
class ExtractedText:
    text: str
    pages: int | None
    truncated: bool


def detect_type(data: bytes) -> DocumentType | None:
    if data[:5] == b"%PDF-":
        return DocumentType.PDF
    if data[:4] == b"PK\x03\x04":
        try:
            with zipfile.ZipFile(io.BytesIO(data)) as archive:
                names = set(archive.namelist())
        except zipfile.BadZipFile:
            return None
        if "word/document.xml" in names and "[Content_Types].xml" in names:
            return DocumentType.DOCX
    return None


def extract_text(data: bytes, doc_type: DocumentType, *, max_pages: int = 20) -> ExtractedText:
    if doc_type is DocumentType.PDF:
        raw, pages, truncated = _extract_pdf(data, max_pages)
    else:
        raw, pages, truncated = _extract_docx(data), None, False
    text = normalize_text(raw)
    if len(text) > MAX_TEXT_CHARS:
        text, truncated = text[:MAX_TEXT_CHARS], True
    return ExtractedText(text=text, pages=pages, truncated=truncated)


def _extract_pdf(data: bytes, max_pages: int) -> tuple[str, int, bool]:
    from pypdf import PdfReader
    from pypdf.errors import PdfReadError

    try:
        reader = PdfReader(io.BytesIO(data))
        if reader.is_encrypted and not reader.decrypt(""):
            raise DocumentError(
                "This PDF is password-protected. Please upload a version without a password.",
                code="document_encrypted",
            )
        total = len(reader.pages)
        parts = [page.extract_text() or "" for page in reader.pages[:max_pages]]
    except DocumentError:
        raise
    except (PdfReadError, ValueError, KeyError, TypeError, OSError) as exc:
        raise DocumentError(
            "This PDF appears to be damaged and can't be read.", code="document_corrupt"
        ) from exc
    return "\n".join(parts), total, total > max_pages


def _extract_docx(data: bytes) -> str:
    from docx import Document

    try:
        with zipfile.ZipFile(io.BytesIO(data)) as archive:
            infos = archive.infolist()
            if (
                len(infos) > MAX_DOCX_ENTRIES
                or sum(i.file_size for i in infos) > MAX_DOCX_UNCOMPRESSED
            ):
                raise DocumentError(
                    "This Word document is too large to process.", code="document_too_large"
                )
        document = Document(io.BytesIO(data))
    except DocumentError:
        raise
    except (zipfile.BadZipFile, KeyError, ValueError, OSError) as exc:
        raise DocumentError(
            "This Word document appears to be damaged and can't be read.", code="document_corrupt"
        ) from exc

    lines = [p.text for p in document.paragraphs]
    # Resumes often lay out sections in tables.
    for table in document.tables:
        for row in table.rows:
            cells = [c.text.strip() for c in row.cells]
            lines.append(" | ".join(dict.fromkeys(c for c in cells if c)))
    return "\n".join(lines)


_CONTROL = re.compile(r"[\x00-\x08\x0b\x0c\x0e-\x1f\x7f]")
# Tabs, no-break space and the typographic spaces U+2000..U+200B.
_SPACES = re.compile("[ \t" + "".join(map(chr, [0xA0, *range(0x2000, 0x200C)])) + "]+")


def normalize_text(text: str) -> str:
    """NFKC, no control characters, single spaces, at most one blank line in a row."""
    text = unicodedata.normalize("NFKC", text).replace("\r\n", "\n").replace("\r", "\n")
    text = _CONTROL.sub("", text)
    lines = [_SPACES.sub(" ", line).strip() for line in text.split("\n")]
    out: list[str] = []
    for line in lines:
        if line or (out and out[-1]):
            out.append(line)
    return "\n".join(out).strip()
