from __future__ import annotations

import io
import zipfile

import pytest

from app.documents.extraction import (
    DocumentError,
    DocumentType,
    detect_type,
    extract_text,
    normalize_text,
)
from tests.resume_fixtures import make_blank_pdf, make_docx, make_pdf


def test_detects_type_from_content_not_name() -> None:
    assert detect_type(make_pdf()) is DocumentType.PDF
    assert detect_type(make_docx()) is DocumentType.DOCX
    assert detect_type(b"MZ\x90\x00 definitely an exe") is None
    # A ZIP that isn't a Word document.
    buffer = io.BytesIO()
    with zipfile.ZipFile(buffer, "w") as archive:
        archive.writestr("readme.txt", "hi")
    assert detect_type(buffer.getvalue()) is None


def test_extracts_pdf_text() -> None:
    result = extract_text(make_pdf(), DocumentType.PDF)
    assert "Senior Backend Engineer, Acme Fintech" in result.text
    assert "Python - 7 years" in result.text
    assert result.pages == 1
    assert not result.truncated


def test_extracts_docx_text_including_tables() -> None:
    from docx import Document

    document = Document()
    document.add_paragraph("Priya Sharma")
    table = document.add_table(rows=1, cols=2)
    table.cell(0, 0).text = "Python"
    table.cell(0, 1).text = "7 years"
    buffer = io.BytesIO()
    document.save(buffer)

    text = extract_text(buffer.getvalue(), DocumentType.DOCX).text
    assert "Priya Sharma" in text
    assert "Python | 7 years" in text


def test_scanned_pdf_has_no_text() -> None:
    assert extract_text(make_blank_pdf(), DocumentType.PDF).text == ""


def test_page_limit_is_reported() -> None:
    from fpdf import FPDF

    pdf = FPDF()
    for i in range(5):
        pdf.add_page()
        pdf.set_font("Helvetica", size=11)
        pdf.cell(0, 10, f"Page {i + 1}")
    result = extract_text(bytes(pdf.output()), DocumentType.PDF, max_pages=2)
    assert "Page 2" in result.text and "Page 3" not in result.text
    assert result.pages == 5 and result.truncated


def test_damaged_pdf_is_a_friendly_error() -> None:
    with pytest.raises(DocumentError) as excinfo:
        extract_text(b"%PDF-1.7\nthis is not really a pdf", DocumentType.PDF)
    assert excinfo.value.code == "document_corrupt"


def test_docx_zip_bomb_is_refused() -> None:
    buffer = io.BytesIO()
    with zipfile.ZipFile(buffer, "w", zipfile.ZIP_DEFLATED) as archive:
        archive.writestr("[Content_Types].xml", "<Types/>")
        archive.writestr("word/document.xml", "<w/>")
        archive.writestr("word/bomb.bin", b"\0" * (60 * 1024 * 1024))  # compresses to ~60 KB
    with pytest.raises(DocumentError) as excinfo:
        extract_text(buffer.getvalue(), DocumentType.DOCX)
    assert excinfo.value.code == "document_too_large"


def test_normalize_text() -> None:
    raw = "Name" + chr(0xA0) + chr(0x2003) + "Surname\r\n\r\n\r\n\x00Skills:\tPython  \n"
    assert normalize_text(raw) == "Name Surname\n\nSkills: Python"
