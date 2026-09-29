"""Real PDF/DOCX documents and a scripted LLM for resume tests."""

from __future__ import annotations

import io
from typing import Any

from app.ai.jobs.analyzer import JobAnalysisExtraction
from app.ai.llm.base import LLMError, T
from app.ai.resume.extraction_schema import ResumeExtraction, from_parsed
from app.ai.resume.schema import ParsedResume

RESUME_LINES = [
    "Priya Sharma",
    "Senior Python Engineer",
    "priya.sharma@example.com | +91 98765 43210 | Bengaluru",
    "linkedin.com/in/priya-sharma | github.com/priya-demo",
    "SUMMARY",
    "Backend engineer with 7 years of experience building Python and Django services.",
    "EXPERIENCE",
    "Senior Backend Engineer, Acme Fintech",
    "Jun 2021 - Present",
    "Built payment APIs in Python and PostgreSQL.",
    "Software Engineer, Initech",
    "Jul 2018 - May 2021",
    "EDUCATION",
    "B.Tech in Computer Science, NIT Trichy, 2014 - 2018",
    "SKILLS",
    "Python - 7 years",
    "Django - 5 years",
    "AWS",
    "CERTIFICATIONS",
    "AWS Certified Developer - Associate (2022)",
]
RESUME_TEXT = "\n".join(RESUME_LINES)


def make_pdf(lines: list[str] = RESUME_LINES) -> bytes:
    from datetime import UTC, datetime

    from fpdf import FPDF

    pdf = FPDF()
    # Fixed timestamp: identical input must give byte-identical files (duplicate checks).
    pdf.set_creation_date(datetime(2026, 1, 1, tzinfo=UTC))
    pdf.add_page()
    pdf.set_font("Helvetica", size=11)
    for line in lines:
        pdf.cell(0, 7, line, new_x="LMARGIN", new_y="NEXT")
    return bytes(pdf.output())


def make_docx(lines: list[str] = RESUME_LINES) -> bytes:
    from docx import Document

    document = Document()
    for line in lines:
        document.add_paragraph(line)
    buffer = io.BytesIO()
    document.save(buffer)
    return buffer.getvalue()


def make_blank_pdf() -> bytes:
    """A valid PDF with no text layer (like a scanned image)."""
    from fpdf import FPDF

    pdf = FPDF()
    pdf.add_page()
    pdf.rect(20, 20, 100, 60)
    return bytes(pdf.output())


SUMMARY = RESUME_LINES[5]


def parsed_resume(**overrides: Any) -> ParsedResume:
    """What a well-behaved model returns for RESUME_TEXT."""
    data: dict[str, Any] = {
        "full_name": "Priya Sharma",
        "email": "priya.sharma@example.com",
        "phone": "+91 98765 43210",
        "location": "Bengaluru",
        "headline": "Senior Python Engineer",
        "summary": SUMMARY,
        "total_years_experience": 7,
        "total_years_evidence": SUMMARY,
        "links": [
            {"kind": "linkedin", "url": "https://linkedin.com/in/priya-sharma"},
            {"kind": "github", "url": "https://github.com/priya-demo"},
        ],
        "skills": [
            {"name": "Python", "years": 7, "evidence": "Python - 7 years"},
            {"name": "Django", "years": 5, "evidence": "Django - 5 years"},
            {"name": "AWS", "years": None, "evidence": "AWS"},
        ],
        "experiences": [
            {
                "title": "Senior Backend Engineer",
                "company": "Acme Fintech",
                "location": None,
                "start": "2021-06",
                "end": None,
                "is_current": True,
                "description": "Built payment APIs in Python and PostgreSQL.",
                "technologies": ["Python", "PostgreSQL"],
                "evidence": "Senior Backend Engineer, Acme Fintech",
            },
            {
                "title": "Software Engineer",
                "company": "Initech",
                "location": None,
                "start": "2018-07",
                "end": "2021-05",
                "is_current": False,
                "description": None,
                "technologies": [],
                "evidence": "Software Engineer, Initech",
            },
        ],
        "education": [
            {
                "institution": "NIT Trichy",
                "degree": "B.Tech",
                "field_of_study": "Computer Science",
                "start_year": 2014,
                "end_year": 2018,
                "grade": None,
                "evidence": "B.Tech in Computer Science, NIT Trichy, 2014 - 2018",
            }
        ],
        "certifications": [
            {
                "name": "AWS Certified Developer - Associate",
                "issuer": None,
                "year": 2022,
                "evidence": "AWS Certified Developer - Associate (2022)",
            }
        ],
    }
    data.update(overrides)
    return ParsedResume.model_validate(data)


class ScriptedLLM:
    """Returns a fixed structured result (or raises), and records what it was sent."""

    name = "scripted"
    model = "test-model"

    def __init__(self, result: ParsedResume | None = None, error: LLMError | None = None) -> None:
        self.result = result or parsed_resume()
        self.error = error
        self.job_analysis: JobAnalysisExtraction | None = None
        self.prompts: list[str] = []

    async def generate(self, *, system: str, prompt: str, max_tokens: int = 4096) -> str:
        return ""

    async def generate_structured(
        self, *, system: str, prompt: str, output_model: type[T], max_tokens: int = 16000
    ) -> T:
        self.prompts.append(prompt)
        if self.error:
            raise self.error
        if output_model is JobAnalysisExtraction:
            from tests.job_fixtures import job_analysis

            return self.job_analysis or job_analysis()  # type: ignore[return-value]
        if output_model is ResumeExtraction:  # what the real parser asks for
            return from_parsed(self.result)  # type: ignore[return-value]
        return self.result  # type: ignore[return-value]

    async def embed(self, texts: list[str]) -> list[list[float]]:
        return [[0.0] for _ in texts]
